"""
One printable script per member (same words as the short deck's speaker notes, plus stage cues).

    python make_member_scripts.py   -> slides/Script_Speaker1.docx, _Speaker2.docx, _Speaker3.docx
"""
from docx import Document
from docx.shared import Pt, RGBColor

from make_deck_short import SCRIPT

CUES = {
    1: "[Slide 1 is already on screen. Smile, look at the TA, and start.]",
    2: "[Click to slide 2. Point at the picture of the junction.]",
    3: "[Click to slide 3. Point at the formula when you say “the formula says”.]",
    4: "[Take the clicker. Click to slide 4.]",
    5: "[Click to slide 5. Point at the table row by row: FCN, CNN, Attention.]",
    6: "[Click to slide 6. Point at the 72 % row, then at the picture with the lines.]",
    7: "[Take the clicker. Click to slide 7. Point at the three big numbers.]",
    8: "[Click to slide 8.]",
    9: "[Click to slide 9. Click the video to play it (or run the live demo). Let it play ~20 s while you speak.]",
    10: "[Click to slide 10. After “Any questions?”, wait. For questions use backup slides B1 and B2.]",
}
PARTS = {1: ("Speaker 1", "Slides 1-3: introduction, the problem, how the car learns"),
         2: ("Speaker 2", "Slides 4-6: DQN, the paper's idea, attention"),
         3: ("Speaker 3", "Slides 7-10: results, our extensions, demo, summary")}

for spk, (title, sub) in PARTS.items():
    doc = Document()
    doc.styles["Normal"].font.name = "Segoe UI"
    doc.styles["Normal"].font.size = Pt(15)
    doc.add_heading(f"{title} - your script", 0)
    secs = sum(v[1] for v in SCRIPT.values() if v[0] == spk)
    doc.add_paragraph(f"{sub}.  About {secs // 60} min {secs % 60} s. Speak slowly, pause at full stops, and look up "
                      "at the TA between sentences - you don't have to read every word exactly.")
    for no in sorted(SCRIPT):
        s, sec, words = SCRIPT[no]
        if s != spk:
            continue
        h = doc.add_paragraph().add_run(f"Slide {no}")
        h.bold, h.font.size, h.font.color.rgb = True, Pt(17), RGBColor(42, 120, 214)
        cue = doc.add_paragraph().add_run(CUES[no])
        cue.italic, cue.font.size, cue.font.color.rgb = True, Pt(12), RGBColor(110, 110, 110)
        doc.add_paragraph(words)
    if spk < 3:
        doc.add_paragraph().add_run(f"[Hand over to Speaker {spk + 1}.]").italic = True
    doc.save(f"slides/Script_Speaker{spk}.docx")
    print(f"-> slides/Script_Speaker{spk}.docx ({secs} s)")
