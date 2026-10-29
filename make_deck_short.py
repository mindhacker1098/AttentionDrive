"""
SHORT and EASY version of the presentation: 10 slides, ~5 minutes of speech, simple words, 3 speakers.

    python make_deck_short.py

-> slides/BanditLAB_Short_Presentation.pptx  (script in each slide's speaker notes)
-> slides/BanditLAB_Short_Script.docx        (printable script)
Backup slides B1-B2 after "Thank you" are only for questions.
"""
import os

from docx import Document
from docx.shared import Pt as DPt, RGBColor as DRGB
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from make_deck import ACCENT, INK, INK_2, SPEAKER_COLORS, SURFACE, bullets, eq, picture, table, text

L, CW = Inches(0.6), Inches(12.1)

# slide -> (speaker, seconds, words to say)
SCRIPT = {
    1: (1, 25, "Good afternoon, everyone. We're Bandit LAB. Have you ever stood at a busy crossing, trying to decide "
               "when it's safe to go? That's exactly the problem in our paper, “Social Attention for Autonomous "
               "Decision-Making in Dense Traffic”, from NeurIPS 2019. It teaches a self-driving car to make that "
               "decision."),
    2: (1, 28, "Here's the situation. Our car wants to turn left, and there's no traffic light. Cars are coming from "
               "every side. If it goes too early, it crashes. If it waits too long, it just sits there forever. We can't "
               "write a rule for every possible case, so instead we let the car learn from experience, by trial and "
               "error. That's reinforcement learning."),
    3: (1, 40, "So how does it learn? Every second, the car looks around: where it is, how fast it's going, and the same "
               "for up to fourteen cars nearby. Then it makes a simple choice: slow down, keep its speed, or speed up. If "
               "it drives smoothly and quickly, it gets a reward. If it crashes, it loses five points. Over time, the car "
               "learns a score, called Q, for each choice. This formula simply says: a good choice now gives a good "
               "reward now and leads to a good situation next. Now I'll hand over to my teammate, who'll explain how the "
               "car actually learns this score."),
    4: (2, 38, "Thanks. Now, there are far too many situations to write all the scores in a table. So we use a neural "
               "network to estimate them. This is called DQN, a Deep Q-Network. To keep the learning stable, it uses "
               "three simple tricks. First, it remembers past moments and studies them in random order, a bit like "
               "revising old notes. Second, it checks its answers against a slowly updated copy of itself, so the target "
               "doesn't keep moving. And third, at the start it tries random moves, just to explore."),
    5: (2, 35, "Now here's the real question the paper asks: how should the car look at the other cars? It compares "
               "three ideas. FCN reads all the cars as one long list of numbers, but the list gets shuffled, so it can't "
               "tell which car is which. CNN looks at a small, blurry picture of the road from above. And attention, the "
               "paper's idea, looks at every car and asks: how important is this one for me? To keep it fair, all three "
               "networks are about the same size."),
    6: (2, 42, "Attention is actually very close to how we drive. Think about it: when you cross a road, you don't look "
               "at every car equally. You focus on the one coming fast towards you. Our car does the same. It gives each "
               "car a danger score, and softmax turns those scores into percentages. Here, the fast car from the left "
               "gets seventy-two percent of the attention, and the rest get very little. And the nice part is, we can "
               "actually see what it's looking at, like the lines in this picture. Now my teammate will show you how "
               "well it works."),
    7: (3, 32, "Thanks. So, did it work? We trained each network three times and tested them on a hundred situations "
               "they had never seen. And just like in the paper, attention came out on top: sixty-seven percent success, "
               "and only twelve percent crashes. FCN was too reckless; it crashed forty-one percent of the time. And CNN "
               "was too careful; it often just waited until the time ran out."),
    8: (3, 38, "But we didn't want to stop there, so we pushed the idea further than the paper. First, we trained one "
               "single agent to handle all twelve routes: left, straight and right, from every road. Attention cut the "
               "crashes in half compared with FCN. Then we built a small city, six junctions joined together, with a "
               "finish line at the end. Even without any extra training, attention finished fifty-two percent of the "
               "journeys, while FCN managed only thirty-four."),
    9: (3, 20, "Let's see it in action. The green band is the route it has to follow. The coloured lines show which cars "
               "it's paying attention to right now. And on the right, you can see what it decides to do, second by "
               "second."),
    10: (3, 22, "So, to wrap up. The big lesson is this: how a car looks at the traffic around it matters just as much "
                "as how it learns. Attention makes it safer, it works with any number of cars, and we can actually "
                "understand its decisions. Thank you very much. We're happy to take any questions."),
}


def slide(prs, no, title, kicker=""):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = SURFACE
    if kicker:
        text(s, kicker, L, Inches(0.35), Inches(9), Inches(0.4), 15, ACCENT, bold=True)
    if title:
        text(s, title, L, Inches(0.72), Inches(12.2), Inches(0.9), 32, INK, bold=True)
    if no in SCRIPT:
        spk, secs, words = SCRIPT[no]
        text(s, f"Speaker {spk}", Inches(11.3), Inches(0.32), Inches(1.6), Inches(0.4), 14,
             SPEAKER_COLORS[spk], bold=True, align=PP_ALIGN.RIGHT)
        s.notes_slide.notes_text_frame.text = f"SPEAKER {spk}  (~{secs} s)\n\n{words}"
    return s


def stat(s, x, y, big, small, color=ACCENT, w=Inches(3.6), size=46):
    text(s, big, x, y, w, Inches(1.0), size, color, bold=True)
    text(s, small, x, y + Inches(0.95), w, Inches(0.9), 16, INK_2)


def build():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)

    s = slide(prs, 1, "", "CS6109 Reinforcement Learning")
    text(s, "Teaching a Self-Driving Car\nto Cross a Busy Junction", L, Inches(1.3), Inches(8.3), Inches(2), 40, INK, True)
    text(s, "Paper: “Social Attention for Autonomous Decision-Making in Dense Traffic”\n"
            "Leurent & Mercat (INRIA + Renault), NeurIPS 2019 workshop", L, Inches(3.4), Inches(8.3), Inches(1), 17, INK_2)
    text(s, "Bandit LAB\nShakti Prasad Nahak (2611AI46)  ·  Akash Prasad (2611AI08)\n"
            "Gourav Rajak (2611AI07)  ·  Dhruv Kumar Singh (2611AI15)", L, Inches(5.0), Inches(8), Inches(1.6), 16, INK)
    picture(s, "slides/frame_title.png", Inches(9.1), Inches(1.4), h=Inches(4.0))

    s = slide(prs, 2, "The problem", "1 · Problem")
    bullets(s, ["Turn left, no traffic light", "Too early → crash", "Too late → never gets through",
                "Too many cases for hand-written rules", "→ learn by trial and error = Reinforcement Learning"],
            L, Inches(1.9), Inches(6.2), Inches(4.5), 24)
    picture(s, "slides/frame_problem.png", Inches(7.2), Inches(1.6), h=Inches(5.6))

    s = slide(prs, 3, "How the car learns", "2 · Reinforcement learning")
    table(s, [["Every second the car…", "In our task"], ["Sees (state)", "position + speed of itself and up to 14 cars"],
              ["Does (action)", "SLOWER · IDLE · FASTER, every second"],
              ["Gets (reward)", "+ for driving fast,  −5 for a crash"]],
          L, Inches(1.8), Inches(9.5), [Inches(3.4), Inches(6.1)], size=19, row_h=Inches(0.62))
    text(s, "Score of an action = reward now + best score next moment:", L, Inches(4.5), Inches(11), Inches(0.5), 19,
         INK_2, True)
    s.shapes.add_picture(eq("bellman_simple", r"Q(s,a)=r+\gamma\,\max_{a'}Q(s',a')", 30), Inches(1.0), Inches(5.1),
                         height=Inches(0.9))
    text(s, "γ = 0.95: how much the future counts", Inches(1.0), Inches(6.15), Inches(6), Inches(0.5), 17, INK_2)

    s = slide(prs, 4, "DQN: a neural network learns the score", "2 · Reinforcement learning")
    bullets(s, ["Too many situations for a table → a neural network predicts Q",
                "Trick 1 - memory: learn from old moments in random order",
                "Trick 2 - a slowly updated copy of the network gives steady targets",
                "Trick 3 - explore: random actions at first (100 % → 5 %)"],
            L, Inches(1.9), CW, Inches(4), 24)

    s = slide(prs, 5, "The paper's idea: how to look at the other cars?", "3 · The paper")
    table(s, [["Network", "How it looks at the cars", "Problem"],
              ["FCN", "one long list of numbers", "can't tell which car is which"],
              ["CNN", "small picture of the road from above", "blurry, slow to decide"],
              ["Attention (paper)", "looks at each car: how important is it?", "-"]],
          L, Inches(1.9), CW, [Inches(3.0), Inches(5.0), Inches(4.1)], size=20, row_h=Inches(0.72))
    text(s, "All three: same learning method (DQN), about 30 000 weights → a fair test", L, Inches(5.2), CW,
         Inches(0.6), 19, INK_2)

    s = slide(prs, 6, "Attention in one example", "3 · The paper")
    table(s, [["Car", "Danger score", "Attention"],
              ["fast car from the left", "2.0", "72 %"], ["far car", "0.5", "16 %"],
              ["car leaving", "0.0", "10 %"], ["ourselves", "−1.5", "2 %"]],
          L, Inches(1.8), Inches(6.6), [Inches(3.4), Inches(1.6), Inches(1.6)], size=19, row_h=Inches(0.6))
    s.shapes.add_picture(eq("softmax_simple", r"\mathrm{attention}_i=\frac{e^{\mathrm{score}_i}}{\sum_j e^{\mathrm{score}_j}}",
                            28), L, Inches(5.2), height=Inches(1.1))
    bullets(s, ["Focus on the dangerous car", "Order of cars does not matter", "We can SEE what it watches"],
            Inches(7.6), Inches(1.9), Inches(5.3), Inches(2.4), 21)
    picture(s, "slides/frame_multi.png", Inches(7.6), Inches(4.1), h=Inches(3.0))

    s = slide(prs, 7, "Result: attention wins, like in the paper", "4 · Results")
    stat(s, L, Inches(1.7), "67 %", "Attention: success\n(only 12 % crashes)", ACCENT)
    stat(s, Inches(4.6), Inches(1.7), "41 %", "FCN: crashes\n(drives too fast)", RGBColor(235, 104, 52))
    stat(s, Inches(8.6), Inches(1.7), "34 %", "CNN: too slow\n(too careful)", RGBColor(27, 175, 122))
    picture(s, "plots/eval_normal.png", Inches(1.3), Inches(3.9), h=Inches(3.4))

    s = slide(prs, 8, "We went further than the paper", "5 · Our extensions")
    text(s, "1.  One agent, all 12 routes", L, Inches(1.8), Inches(6), Inches(0.6), 24, INK, True)
    stat(s, L, Inches(2.4), "18 % vs 34 %", "crashes: attention vs FCN (crashes halved)", ACCENT, Inches(6.2), 38)
    text(s, "2.  A city of 6 junctions to a finish", L, Inches(4.4), Inches(6.2), Inches(0.6), 24, INK, True)
    stat(s, L, Inches(5.0), "52 % vs 34 %", "journeys finished: attention vs FCN (no extra training)", ACCENT,
         Inches(6.2), 38)
    picture(s, "slides/frame_journey.png", Inches(6.9), Inches(1.7), w=Inches(6.0))

    s = slide(prs, 9, "Demo", "6 · Demo")
    if os.path.exists("videos/journey.mp4"):
        s.shapes.add_movie("videos/journey.mp4", L, Inches(1.5), Inches(9.1), Inches(5.83),
                           poster_frame_image="slides/frame_journey.png", mime_type="video/mp4")
    bullets(s, ["Green band = route", "Lines = attention", "Right panel = decision"],
            Inches(9.9), Inches(1.8), Inches(3.2), Inches(3), 20)

    s = slide(prs, 10, "Summary", "7 · Takeaway")
    bullets(s, ["How the network looks at other cars matters as much as the learning method",
                "Attention: safer, any number of cars, easy to understand",
                "Reproduced the paper + extended it to 12 routes and a city",
                "Limits: simulator only, speed control only (no lane changes)"],
            L, Inches(1.9), CW, Inches(3.6), 24)
    text(s, "Thank you!  Questions?", L, Inches(5.6), CW, Inches(1), 36, ACCENT, True)

    s = slide(prs, 101, "Backup B1 · All results", "For questions")
    table(s, [["Test", "Attention", "FCN", "CNN", "Random"],
              ["Left turn: success / crash %", "67 / 12", "55 / 41", "48 / 17", "20 / 40"],
              ["Heavy traffic: success / crash %", "42 / 17", "32 / 61", "27 / 25", "15 / 54"],
              ["12 routes: success / crash %", "70 / 18", "65 / 34", "-", "39 / 28"],
              ["City journey: finished / crash %", "52.5 / 45.5", "34 / 65.5", "-", "37 / 43"]],
          L, Inches(1.8), CW, [Inches(4.4), Inches(2.0), Inches(2.0), Inches(1.8), Inches(1.9)], size=17,
          row_h=Inches(0.6))
    text(s, "Fine-tuning on the city made it worse (44 % finished): it learned to drive more aggressively.", L,
         Inches(5.1), CW, Inches(0.6), 17, INK_2)

    s = slide(prs, 102, "Backup B2 · Where things are in the code", "For questions")
    table(s, [["File", "What it does"],
              ["common.py", "the road task + the 3 networks (attention = attend function)"],
              ["train.py", "DQN training: memory, slowly updated copy, exploration"],
              ["evaluate.py", "test on 100 new situations"],
              ["journey.py", "the city of 6 junctions + route planner"],
              ["demo.py", "the live demo with attention lines"]],
          L, Inches(1.8), CW, [Inches(3.0), Inches(9.1)], size=18, row_h=Inches(0.6))

    out = "slides/BanditLAB_Short_Presentation.pptx"
    prs.save(out)
    print("->", out)


def script_docx():
    doc = Document()
    doc.styles["Normal"].font.name = "Segoe UI"
    doc.styles["Normal"].font.size = DPt(13)
    total = sum(v[1] for v in SCRIPT.values())
    doc.add_heading("Bandit LAB - short script (read aloud)", 0)
    doc.add_paragraph(f"10 slides, about {total // 60} min {total % 60} s of speech + the demo. "
                      "Each speaker clicks through their own slides.")
    names = {1: "Speaker 1 - slides 1-3", 2: "Speaker 2 - slides 4-6", 3: "Speaker 3 - slides 7-10"}
    current = None
    for no in sorted(SCRIPT):
        spk, secs, words = SCRIPT[no]
        if spk != current:
            current = spk
            doc.add_heading(names[spk], 1)
        r = doc.add_paragraph().add_run(f"Slide {no}")
        r.bold = True
        r.font.color.rgb = DRGB(42, 120, 214)
        doc.add_paragraph(words)
    doc.add_heading("If the TA asks", 1)
    for q, a in [("What is DQN?", "It's Q-learning, but instead of a table, a neural network guesses the score of "
                                  "each action."),
                 ("What is attention?", "The car gives each other car an importance score, and softmax turns those "
                                        "scores into percentages that add up to one."),
                 ("Why doesn't the order of the cars matter?", "Attention just adds up all the cars with their weights, "
                                                               "and a sum is the same in any order."),
                 ("Why only three runs per network?", "Each run takes about an hour on a laptop. The paper used 120, so "
                                                      "we show the uncertainty on our graphs."),
                 ("Why can't it drive around obstacles?", "It only controls its speed, not the steering. Changing lanes "
                                                          "would be a nice next step.")]:
        p = doc.add_paragraph()
        p.add_run(q + "  ").bold = True
        p.add_run(a)
    out = "slides/BanditLAB_Short_Script.docx"
    doc.save(out)
    print("->", out, f"({total} s)")


if __name__ == "__main__":
    build()
    script_docx()
