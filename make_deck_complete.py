"""
COMPLETE presentation (16 slides): the paper (problem, MDP, RL maths, DQN, the three networks, attention maths)
and our work (implementation, results, 12-route extension, city journey, demo), with formulas.
Natural-sounding speaker notes for 3 speakers (~9 min) + a printable script.

    python make_deck_complete.py
-> slides/BanditLAB_Complete_Presentation.pptx
-> slides/BanditLAB_Complete_Script.docx
"""
import os

from docx import Document
from docx.shared import Pt as DPt, RGBColor as DRGB
from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from make_deck import (ACCENT, INK, INK_2, SPEAKER_COLORS, SURFACE, bullets, equation, picture, table, text)

L, CW = Inches(0.6), Inches(12.1)

SCRIPT = {
    1: (1, 20, "Good afternoon, everyone. We're Bandit LAB. Today we'll present our implementation of “Social Attention "
               "for Autonomous Decision-Making in Dense Traffic”, by Leurent and Mercat from INRIA and Renault, "
               "published at a NeurIPS 2019 workshop. We reproduced the paper, and then we extended it in two ways."),
    2: (1, 30, "Picture a car that has to turn left at a busy junction with no traffic light. Go too early, and it "
               "crashes. Wait too long, and it never gets through. That's called the freezing robot problem. Real "
               "drivers handle this by watching only the few cars that matter. Writing rules for every case is "
               "impossible, so the car learns by trial and error, with reinforcement learning."),
    3: (1, 40, "We describe the task as a Markov Decision Process. The state is a table with one row per car: our car "
               "plus up to fourteen others, and seven numbers per row: whether the car is there, its position, its "
               "velocity and its heading. We scale everything to between minus one and one. Every second, the agent "
               "picks slower, idle or faster; the steering is automatic. The reward at the bottom gives up to one point "
               "per second for driving fast, minus five for a crash, and one for reaching the exit."),
    4: (1, 35, "The agent wants to maximise the return: the sum of future rewards, where rewards further away count a "
               "bit less, because of gamma, zero point nine five. The Q-value tells us how good an action is in a given "
               "state. The best Q-values satisfy the Bellman equation: the value now equals the reward now, plus gamma "
               "times the best value in the next state. And once we know Q, driving is easy: just pick the action with "
               "the highest Q."),
    5: (1, 45, "Since the state is continuous, a neural network estimates Q. That's DQN. We train it towards the target "
               "y, the reward plus the discounted value of the next state, by minimising the squared error. Three tricks "
               "keep training stable: a replay buffer of fifteen thousand past moments, sampled at random; a target "
               "network, a frozen copy updated every five hundred and twelve steps; and Double DQN, where one network "
               "picks the next action and the other one scores it. Exploration follows this epsilon schedule, from one "
               "down to five percent. Over to my teammate for the paper's main idea."),
    6: (2, 30, "Thanks. The paper keeps DQN fixed and asks one question: how should the network read the other cars? A "
               "plain list changes size with the number of cars, and it depends on their order. But re-ordering the cars "
               "shouldn't change the decision. That's called permutation invariance, written here as an equation. To "
               "test it, the order of the cars is shuffled at every step, in the paper and in our code."),
    7: (2, 30, "The paper compares three networks of about the same size, around thirty thousand weights. FCN flattens "
               "the table into one long vector, so it isn't order-invariant. CNN uses a thirty-two by thirty-two grid "
               "seen from above; it is order-invariant, but coarse. And Ego-Attention is the paper's model. Our weight "
               "counts match the paper's Table 1 very closely."),
    8: (2, 45, "Here's how ego-attention works. Each car is encoded by a small shared network into a vector e. Our car "
               "creates a query, and every car creates a key and a value. The score of each car is the query dot its key, "
               "divided by the square root of the head size. Softmax turns the scores into weights alpha that add up to "
               "one, and empty rows are masked out. The output is the weighted sum of the values, added back to our "
               "car's own encoding, and a decoder gives the three Q-values. There are two heads, so this happens twice, "
               "in parallel."),
    9: (2, 35, "A quick example. Scores of two, zero point five, zero and minus one point five become weights of "
               "seventy-two, sixteen, ten and two percent. So the network focuses on the dangerous car. Because it's a "
               "weighted sum, shuffling the cars changes nothing, and it works for any number of cars. And we can draw "
               "these weights as lines, like in this picture, to see what the agent is watching."),
    10: (2, 30, "We implemented all of this from scratch in PyTorch, and took the hyper-parameters straight from the "
                "authors' own code. We trained three agents with three seeds each, four thousand episodes per run, and "
                "tested each one on the same hundred unseen scenarios, with no exploration. Now my teammate will take you "
                "through the results."),
    11: (3, 40, "Thanks. Our training curves show the same three behaviours as the paper. FCN is fast but reckless, so it "
                "crashes. CNN is too cautious and slow. Ego-attention gets the highest reward. On the test, attention "
                "succeeds sixty-seven percent of the time with just twelve percent crashes. FCN crashes forty-one percent "
                "of the time, and CNN runs out of time in a third of the cases. With heavier traffic than in training, "
                "the ranking stays the same."),
    12: (3, 35, "Next, we went beyond the paper. The paper trains for one left turn, and suggests giving the agent its "
                "destination as future work. We did that: the car can start on any of four roads and go left, straight "
                "or right, twelve routes in total. We add the direction to each car's destination: for our car it's the "
                "goal, for the others it works like a turn signal. One attention agent learned all twelve routes and "
                "halved the crashes compared with FCN, with the biggest gain on left turns."),
    13: (3, 40, "Our second extension is a long journey through a city of six junctions, to a finish line. A "
                "shortest-path planner chooses the turns, and our agent decides when to go. At each junction we give it "
                "positions relative to that junction, so its skill transfers. Without extra training, attention finished "
                "fifty-two and a half percent of the journeys, against thirty-four for FCN. Long journeys are harder "
                "because the risks multiply, as this formula shows. And honestly, fine-tuning on the city made it worse, "
                "because it learned to be more aggressive."),
    14: (3, 25, "Here's our demo. The green band is the route, the lines show which cars the agent is watching, and the "
                "panel on the right shows its Q-values and its decision, second by second."),
    15: (3, 30, "To conclude: how the network reads the traffic matters as much as the learning algorithm. Attention is "
                "safer, works with any number of cars, and is easy to interpret. Our limitations: it's a simulator, the "
                "other drivers follow simple rules, the car only controls its speed, and we used three seeds instead of "
                "a hundred and twenty. Next steps would be lane changes, more seeds and real traffic data."),
    16: (3, 10, "Thank you very much. We're happy to take your questions."),
}


def slide(prs, no, title, kicker):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = SURFACE
    text(s, kicker, L, Inches(0.32), Inches(9), Inches(0.4), 14, ACCENT, bold=True)
    if title:
        text(s, title, L, Inches(0.68), Inches(12.2), Inches(0.9), 28, INK, bold=True)
    spk, secs, words = SCRIPT[no]
    text(s, f"Speaker {spk}", Inches(11.3), Inches(0.3), Inches(1.6), Inches(0.4), 13, SPEAKER_COLORS[spk],
         bold=True, align=PP_ALIGN.RIGHT)
    text(s, f"{no}", Inches(12.5), Inches(7.0), Inches(0.6), Inches(0.4), 11, INK_2, align=PP_ALIGN.RIGHT)
    s.notes_slide.notes_text_frame.text = f"SPEAKER {spk}  (~{secs} s)\n\n{words}"
    return s


def build():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)

    # 1 title
    s = slide(prs, 1, "", "CS6109 Reinforcement Learning  ·  Paper implementation")
    text(s, "Teaching a Self-Driving Car\nto Cross a Busy Intersection\nwith Attention", L, Inches(1.0), Inches(8.4),
         Inches(2.2), 34, INK, True)
    text(s, "Paper: E. Leurent & J. Mercat (INRIA + Renault), “Social Attention for Autonomous Decision-Making in Dense "
            "Traffic”, ML for Autonomous Driving Workshop, NeurIPS 2019 (arXiv 1911.12250)", L, Inches(3.45),
         Inches(8.4), Inches(1), 16, INK_2)
    text(s, "Reproduction  +  12-route extension  +  long journey through a city", L, Inches(4.45), Inches(8.4),
         Inches(0.5), 18, ACCENT, True)
    text(s, "Bandit LAB\nShakti Prasad Nahak (2611AI46)  ·  Akash Prasad (2611AI08)\n"
            "Gourav Rajak (2611AI07)  ·  Dhruv Kumar Singh (2611AI15)", L, Inches(5.25), Inches(8), Inches(1.6), 16, INK)
    picture(s, "slides/frame_title.png", Inches(9.2), Inches(1.5), h=Inches(3.9))

    # 2 problem
    s = slide(prs, 2, "The problem: an unprotected left turn", "1 · Problem and motivation")
    bullets(s, ["No traffic light, cars arriving from every direction",
                "Too early → collision;  too late → never gets through (“freezing robot”)",
                "Human drivers watch only the few cars that matter",
                "Rule-based planners (finite-state machines) do not scale to such interactions",
                "→ learn the behaviour by trial and error: Reinforcement Learning",
                "Paper's focus: HOW the network should represent the surrounding cars"],
            L, Inches(1.75), Inches(6.6), Inches(5), 18)
    picture(s, "slides/frame_problem.png", Inches(7.4), Inches(1.6), h=Inches(5.5))

    # 3 MDP
    s = slide(prs, 3, "Formulation as a Markov Decision Process (S, A, P, R, γ)", "2 · Reinforcement learning")
    table(s, [["Element", "In this task"],
              ["State s", "15 cars × 7 features: presence, x, y, vx, vy, cos θ, sin θ (row 0 = our car); "
                          "x, y ÷ 100 and v ÷ 20 → values in [−1, 1]; other cars' order shuffled every step"],
              ["Action a", "every 1 s: SLOWER · IDLE · FASTER (target speed 0 / 4.5 / 9 m/s); steering is automatic"],
              ["Transition P", "highway-env simulator, 15 Hz physics; other drivers follow the IDM model with road priorities"],
              ["Episode", "ends on crash, on arrival, or after 13 s"],
              ["Discount γ", "0.95  →  effective horizon ≈ 1 / (1 − γ) = 20 decisions"]],
          L, Inches(1.65), CW, [Inches(2.2), Inches(9.9)], size=15, row_h=Inches(0.55))
    text(s, "Reward per step:", L, Inches(5.35), Inches(4), Inches(0.4), 16, INK_2, True)
    equation(s, "reward", L, Inches(5.8), h=Inches(0.62))

    # 4 RL maths
    s = slide(prs, 4, "Return, Q-value and the Bellman equation", "2 · Reinforcement learning")
    for i, (label, name, hgt) in enumerate([("Return - discounted sum of future rewards:", "return", 0.85),
                                            ("Action value - expected return after action a in state s:", "qdef", 0.7),
                                            ("Bellman optimality equation (paper, Section 2):", "bellman", 0.7),
                                            ("Greedy policy:", "policy", 0.7)]):
        y = Inches(1.6 + 1.37 * i)
        text(s, label, L, y, Inches(11), Inches(0.4), 16, INK_2, True)
        equation(s, name, Inches(1.0), y + Inches(0.42), h=Inches(hgt))

    # 5 DQN
    s = slide(prs, 5, "Deep Q-Network (DQN) with Double-DQN targets", "2 · Reinforcement learning")
    text(s, "Target and loss (θ = online network, θ̄ = target network, d = 1 at crash / arrival):", L, Inches(1.55),
         CW, Inches(0.4), 16, INK_2, True)
    equation(s, "target", L, Inches(2.0), h=Inches(0.65))
    equation(s, "loss", L, Inches(2.75), h=Inches(0.8))
    bullets(s, ["Replay buffer of 15 000 transitions, random mini-batches of 64 → breaks correlation",
                "Target network θ̄ copied every 512 steps → stable regression target",
                "Double DQN: θ selects a′, θ̄ evaluates it → reduces over-estimation of Q",
                "Adam (lr 5·10⁻⁴), gradients clipped to [−1, 1], one update per environment step"],
            L, Inches(3.7), CW, Inches(2.4), 17)
    text(s, "ε-greedy exploration:", L, Inches(6.25), Inches(3), Inches(0.5), 16, INK_2, True)
    equation(s, "eps", Inches(3.3), Inches(6.2), h=Inches(0.5))

    # 6 representation
    s = slide(prs, 6, "The paper's question: how to represent the other cars?", "3 · The paper's idea")
    bullets(s, ["DQN stays the same - only the network that reads the state changes",
                "List of features: compact and accurate, but its size varies and it depends on the ordering",
                "Spatial grid: order-free, but trades accuracy against size",
                "A driving policy should not depend on the order in which cars are listed:"],
            L, Inches(1.75), CW, Inches(3), 19)
    equation(s, "perm", Inches(1.2), Inches(4.3), h=Inches(0.65))
    text(s, "τ = any re-ordering of cars 1…N  →  “permutation invariance” (paper eq. 2). "
            "We shuffle the car order every step to test it.", Inches(1.2), Inches(5.15), Inches(11), Inches(0.8), 16, INK_2)

    # 7 three networks
    s = slide(prs, 7, "Three architectures, same DQN, ≈ 30 k weights each", "3 · The paper's idea")
    table(s, [["Agent", "Input", "Layers", "Weights (paper / ours)", "Order-free?"],
              ["FCN/List", "15 × 7 table, flattened (105)", "105 → 128 → 128 → 3", "3.0e4 / 30 467", "no"],
              ["CNN/Grid", "32 × 32 grid, 2 m cells, 7 channels", "3 conv (k2, s2): 16-32-64 → 20 → 3", "3.2e4 / 31 363", "yes"],
              ["Ego-Attention", "15 × 7 table, any order / size", "encoders 7→64→64, 2 heads, decoder 64→64→64→3",
               "3.4e4 / 34 243", "yes"]],
          L, Inches(1.8), CW, [Inches(2.0), Inches(3.0), Inches(4.0), Inches(1.9), Inches(1.2)], size=15,
          row_h=Inches(0.65))
    bullets(s, ["Same number of weights → the comparison isolates the effect of the representation",
                "Paper's training: 4000 episodes, 120 seeds;  ours: 4000 episodes, 3 seeds"],
            L, Inches(4.7), CW, Inches(1.5), 17)

    # 8 attention maths
    s = slide(prs, 8, "Ego-attention: the equations", "3 · The paper's idea")
    for i, (label, name, hgt) in enumerate([
            ("1. Shared encoders; our car (ego) emits the query, every car a key and a value:", "enc", 0.5),
            ("2. Scaled dot-product attention, empty rows masked → weights sum to 1:", "alpha", 0.95),
            ("3. Weighted sum, residual with the ego encoding, decoder → Q-values:", "out", 0.8)]):
        y = Inches(1.6 + 1.6 * i)
        text(s, label, L, y, Inches(12), Inches(0.4), 16, INK_2, True)
        equation(s, name, Inches(1.0), y + Inches(0.45), h=Inches(hgt))
    text(s, "d = 64, 2 heads of d_h = 32, computed in parallel and concatenated before W_O.", L, Inches(6.5), CW,
         Inches(0.5), 15, INK_2)

    # 9 example + interpretability
    s = slide(prs, 9, "Attention in one example, and why it works", "3 · The paper's idea")
    table(s, [["Car", "q·k / √d_h", "α (softmax)"],
              ["fast car from the left", "2.0", "0.72"], ["far car on the right", "0.5", "0.16"],
              ["car already leaving", "0.0", "0.10"], ["ourselves", "−1.5", "0.02"]],
          L, Inches(1.7), Inches(6.3), [Inches(3.2), Inches(1.5), Inches(1.6)], size=16, row_h=Inches(0.5))
    bullets(s, ["Focus on the car that matters (72 %)",
                "Weighted sum → order of cars irrelevant",
                "Any number of cars (masked rows get 0)",
                "Interpretable: weights drawn as lines →"],
            L, Inches(4.5), Inches(6.3), Inches(2.6), 17)
    picture(s, "slides/frame_multi.png", Inches(7.1), Inches(1.7), w=Inches(5.9))

    # 10 implementation
    s = slide(prs, 10, "Our implementation and experimental setup", "4 · Our work")
    table(s, [["Item", "Details"],
              ["Code", "PyTorch from scratch: common.py (env + 3 networks), train.py (DQN), evaluate.py, journey.py, demo.py"],
              ["Hyper-parameters", "from the authors' rl-agents configs: γ 0.95, lr 5·10⁻⁴, batch 64, buffer 15 000, "
                                   "target every 512 steps, ε decay τ = 15 000"],
              ["Training", "3 agents × 3 seeds × 4000 episodes = 9 runs (CPU, ~1 h each in parallel)"],
              ["Evaluation", "100 unseen scenarios (fixed seeds), the same for every agent, greedy policy"],
              ["Metrics", "success (arrived) %, crash %, timeout %, average speed, return"],
              ["Our experiment", "denser traffic than in training: 15 starting cars (vs 10), 90 % spawn rate (vs 60 %)"]],
          L, Inches(1.65), CW, [Inches(2.4), Inches(9.7)], size=15, row_h=Inches(0.6))

    # 11 results
    s = slide(prs, 11, "Reproduction results: the paper's ranking holds", "5 · Results")
    picture(s, "plots/training_curves.png", L, Inches(1.5), w=Inches(7.6))
    table(s, [["Test (100 scenarios)", "Success %", "Crash %"],
              ["Random", "20", "40"], ["FCN/List", "55", "41"], ["CNN/Grid", "48", "17"], ["Ego-Attention", "67", "12"]],
          Inches(8.45), Inches(1.6), Inches(4.5), [Inches(2.3), Inches(1.1), Inches(1.1)], size=15, row_h=Inches(0.45))
    bullets(s, ["FCN: fast, crashes (“blind”)", "CNN: slow, 34 % timeouts", "Attention: best reward",
                "Dense traffic: attention 42 / 17, FCN 32 / 61"], Inches(8.45), Inches(4.05), Inches(4.6),
            Inches(3), 15)

    # 12 extension 1
    s = slide(prs, 12, "Extension 1: one agent for all 12 routes (goal conditioning)", "6 · Beyond the paper")
    bullets(s, ["Paper: left turn only; its Section 4 proposes adding the destination as a feature",
                "Ours: 4 start roads × {left, straight, right} = 12 routes, one network",
                "Two extra features per car - unit vector towards its destination:"],
            L, Inches(1.65), Inches(6.4), Inches(2.4), 16)
    equation(s, "goal", Inches(1.0), Inches(3.75), h=Inches(0.75))
    bullets(s, ["ego car → its goal;  other cars → their turn signal",
                "2 seeds × 6000 episodes, 150 test scenarios",
                "Attention 70 % success / 18 % crash  vs  FCN 65 % / 34 %",
                "Left turn (hardest): 52 % vs 41 %"],
            L, Inches(4.7), Inches(6.4), Inches(2.6), 16)
    picture(s, "plots/multi_by_turn.png", Inches(7.0), Inches(1.9), w=Inches(6.0))

    # 13 extension 2
    s = slide(prs, 13, "Extension 2: a long journey through a city", "6 · Beyond the paper")
    bullets(s, ["6 intersections (2 × 3 grid), finish line 3–4 junctions away, up to 26 cars",
                "Route planner: breadth-first shortest path chooses the turns",
                "Agent sees each junction as in training (positions relative to it):"],
            L, Inches(1.6), Inches(7.2), Inches(1.9), 16)
    equation(s, "local", Inches(1.0), Inches(3.2), h=Inches(0.7))
    text(s, "c_k = centre of the junction ahead\ngoal = the exit to take there", Inches(3.3), Inches(3.25),
         Inches(4.2), Inches(0.7), 14, INK_2)
    table(s, [["100 journeys, identical traffic", "Finish %", "Crash %"],
              ["Random", "37", "43"], ["FCN, zero-shot", "34", "66"],
              ["Attention, zero-shot", "52.5", "45.5"], ["Attention, fine-tuned", "44", "55.5"]],
          L, Inches(4.1), Inches(7.0), [Inches(4.0), Inches(1.5), Inches(1.5)], size=15, row_h=Inches(0.4))
    equation(s, "compound", L, Inches(6.3), h=Inches(0.62))
    picture(s, "slides/frame_journey.png", Inches(7.9), Inches(1.6), w=Inches(5.1))

    # 14 demo
    s = slide(prs, 14, "Demo: driving to the finish line", "7 · Demo")
    if os.path.exists("videos/journey.mp4"):
        s.shapes.add_movie("videos/journey.mp4", L, Inches(1.5), Inches(9.1), Inches(5.83),
                           poster_frame_image="slides/frame_journey.png", mime_type="video/mp4")
    bullets(s, ["Green band = planned route", "Lines = attention weights (2 heads)", "Panel = Q-values + action",
                "Mini-map + finish line", "Live: python demo.py --run runs/multi_attention_s2 --journey"],
            Inches(9.9), Inches(1.7), Inches(3.2), Inches(5.5), 15)

    # 15 conclusion
    s = slide(prs, 15, "Conclusion, limitations and future work", "8 · Takeaways")
    bullets(s, ["The state representation matters as much as the RL algorithm",
                "Ego-attention: permutation invariant, any number of cars, interpretable",
                "Reproduced the paper: attention 67 % / 12 % vs FCN 55 / 41 and CNN 48 / 17 (success / crash)",
                "Extended it: 12 routes with goal + turn-signal inputs (crashes halved); city journey 52.5 % vs 34 %",
                "Limitations: simulator; rule-based other drivers; speed control only (no lane changes); 3 seeds",
                "Future: lane changes to pass obstacles, more seeds, longer / safer fine-tuning, real traffic data"],
            L, Inches(1.75), CW, Inches(5), 19)

    # 16 references + thanks
    s = slide(prs, 16, "References  ·  Thank you!", "Questions?")
    bullets(s, ["E. Leurent, J. Mercat. Social Attention for Autonomous Decision-Making in Dense Traffic. "
                "NeurIPS 2019 ML4AD workshop, arXiv 1911.12250",
                "E. Leurent. An environment for autonomous driving decision-making (highway-env), 2018",
                "V. Mnih et al. Human-level control through deep reinforcement learning (DQN). Nature, 2015",
                "H. van Hasselt, A. Guez, D. Silver. Deep RL with Double Q-learning. AAAI, 2016",
                "A. Vaswani et al. Attention Is All You Need. NeurIPS, 2017",
                "Authors' code and configs: github.com/eleurent/rl-agents  ·  Our code: rl_highway/"],
            L, Inches(1.75), CW, Inches(4.5), 16)
    text(s, "Thank you - we're happy to take your questions.", L, Inches(6.2), CW, Inches(0.8), 28, ACCENT, True)

    # backup: flow diagrams (python make_flow.py), no script - show them if the TA asks how it all fits together
    for name, label in [("flow_project", "B1 · Project flow"), ("flow_training", "B2 · Training loop"),
                        ("flow_attention", "B3 · One decision with attention")]:
        s = prs.slides.add_slide(prs.slide_layouts[6])
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = SURFACE
        picture(s, f"plots/{name}.png", Inches(0.0), Inches(0.0), w=Inches(13.333))
        text(s, f"Backup  ·  {label}", L, Inches(0.25), Inches(9), Inches(0.4), 13, ACCENT, bold=True)

    out = "slides/BanditLAB_Complete_Presentation.pptx"
    prs.save(out)
    print("->", out)


def script_docx():
    doc = Document()
    doc.styles["Normal"].font.name = "Segoe UI"
    doc.styles["Normal"].font.size = DPt(13)
    total = sum(v[1] for v in SCRIPT.values())
    doc.add_heading("Bandit LAB - script for the complete presentation", 0)
    doc.add_paragraph(f"16 slides, about {total // 60} min {total % 60} s of speech (plus the demo). "
                      "Speak slowly, pause at full stops, and look up at the TA between sentences.")
    names = {1: "Speaker 1 - slides 1-5: problem, MDP, RL maths, DQN",
             2: "Speaker 2 - slides 6-10: the paper's idea, networks, attention, our implementation",
             3: "Speaker 3 - slides 11-16: results, extensions, demo, conclusion"}
    current = None
    for no in sorted(SCRIPT):
        spk, secs, words = SCRIPT[no]
        if spk != current:
            current = spk
            doc.add_heading(names[spk], 1)
        r = doc.add_paragraph().add_run(f"Slide {no}  (~{secs} s)")
        r.bold = True
        r.font.color.rgb = DRGB(42, 120, 214)
        doc.add_paragraph(words)
    out = "slides/BanditLAB_Complete_Script.docx"
    doc.save(out)
    print("->", out, f"({total} s)")


if __name__ == "__main__":
    build()
    script_docx()
